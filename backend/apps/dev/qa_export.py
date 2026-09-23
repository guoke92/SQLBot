"""QA-admin Excel export aligned with exports/export_qa_chats.py."""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlmodel import Session, select

from apps.chat.models.chat_model import ChatRecord
from apps.dev.batch_export import extract_record_sql
from apps.dev.catalog import list_chats
from apps.system.models.system_model import WorkspaceModel

HEADERS = [
    "操作人名",
    "chatid",
    "对话brief",
    "每轮对话问题描述",
    "每轮对话反馈",
    "反馈描述",
    "每轮对话生成的sql语句",
    "对话时间",
    "备注",
]


def _fmt_time(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        text = value.isoformat(sep=" ", timespec="milliseconds")
    else:
        text = str(value).strip()
    if text.endswith("Z"):
        text = text[:-1]
    if "T" in text:
        text = text.replace("T", " ")
    if "." in text:
        head, frac = text.split(".", 1)
        text = f"{head}.{frac[:3]}" if frac[:3].isdigit() else head
    return text


def _feedback_label(record: ChatRecord) -> str:
    fb = str(record.feedback or "").strip().lower()
    return {"up": "认同", "down": "不认同"}.get(fb, "")


def _is_admin_operator(operator_name: str, user_account: str) -> bool:
    name = (operator_name or "").strip().lower()
    account = (user_account or "").strip().lower()
    return name in {"administrator", "admin"} or account in {"admin", "administrator"}


def build_qa_chats_workbook(
    session: Session,
    *,
    oid: int,
    create_by: int | None = None,
    q: str | None = None,
    feedback: str | None = None,
    chat_ids: list[int] | None = None,
    datasource: int | None = None,
    created_from: datetime | None = None,
    created_to: datetime | None = None,
    chat_type: str | None = "chat",
    skip_admin: bool = True,
) -> tuple[bytes, dict[str, Any]]:
    """Build the review Excel workbook. Returns (xlsx_bytes, summary_stats)."""
    chats = list_chats(
        session,
        oid=oid,
        create_by=create_by,
        q=q,
        feedback=feedback,
        chat_ids=chat_ids,
        datasource=datasource,
        created_from=created_from,
        created_to=created_to,
        chat_type=chat_type,
    )
    workspace = session.get(WorkspaceModel, oid)
    workspace_name = workspace.name if workspace is not None else None

    skipped_empty = skipped_admin = skipped_no_sql = 0
    eligible: list[dict[str, Any]] = []
    for chat in chats:
        if int(chat.get("record_count") or 0) <= 0:
            skipped_empty += 1
            continue
        operator_name = str(chat.get("user_name") or chat.get("user_account") or "")
        user_account = str(chat.get("user_account") or "")
        if skip_admin and _is_admin_operator(operator_name, user_account):
            skipped_admin += 1
            continue
        eligible.append(chat)

    by_chat: dict[str, list[dict[str, str]]] = {}
    chat_meta: dict[str, dict[str, str]] = {}

    if eligible:
        ids = [int(item["id"]) for item in eligible]
        records = list(
            session.exec(
                select(ChatRecord)
                .where(ChatRecord.chat_id.in_(ids))
                .order_by(ChatRecord.chat_id, ChatRecord.create_time, ChatRecord.id)
            ).all()
        )
        records_by_chat: dict[int, list[ChatRecord]] = {}
        for record in records:
            if record.chat_id is None:
                continue
            records_by_chat.setdefault(int(record.chat_id), []).append(record)

        for chat in eligible:
            chat_id = str(chat["id"])
            chat_pk = int(chat["id"])
            chat_records = records_by_chat.get(chat_pk) or []
            if not chat_records:
                skipped_empty += 1
                continue

            operator_name = str(
                chat.get("user_name") or chat.get("user_account") or chat.get("create_by") or ""
            )
            brief = str(chat.get("brief") or "")
            chat_time = _fmt_time(chat.get("create_time"))

            all_turns: list[dict[str, str]] = []
            for rec in chat_records:
                question = (rec.question or "").strip()
                if not question:
                    continue
                turn_time = _fmt_time(rec.create_time) or chat_time
                all_turns.append(
                    {
                        "每轮对话问题描述": question,
                        "每轮对话反馈": _feedback_label(rec),
                        "反馈描述": (rec.feedback_comment or "").strip(),
                        "每轮对话生成的sql语句": extract_record_sql(rec),
                        "对话时间": turn_time,
                        "_sort_time": turn_time,
                    }
                )
            all_turns.sort(key=lambda row: row.get("_sort_time") or "")

            with_sql = [
                turn
                for turn in all_turns
                if (turn.get("每轮对话生成的sql语句") or "").strip()
            ]
            without_sql = [
                turn
                for turn in all_turns
                if not (turn.get("每轮对话生成的sql语句") or "").strip()
            ]
            if not with_sql:
                skipped_no_sql += 1
                continue

            note_parts: list[str] = []
            for idx, turn in enumerate(without_sql, 1):
                question = turn["每轮对话问题描述"]
                note_parts.append(f"{idx}. {question}" if len(without_sql) > 1 else question)
            remark = (
                "未生成SQL的提问：\n" + "\n".join(note_parts) if note_parts else ""
            )

            chat_meta[chat_id] = {
                "操作人名": operator_name,
                "chatid": chat_id,
                "对话brief": brief,
                "备注": remark,
            }
            by_chat[chat_id] = with_sql

    ordered_chat_ids = sorted(by_chat.keys(), key=lambda x: int(x) if x.isdigit() else x)
    rows: list[dict[str, str]] = []
    merge_ranges: list[tuple[int, int]] = []
    excel_row = 2
    agree = disagree = 0
    for chat_id in ordered_chat_ids:
        turns = by_chat[chat_id]
        if not turns:
            continue
        meta = chat_meta[chat_id]
        start = excel_row
        for turn in turns:
            fb = turn["每轮对话反馈"]
            if fb == "认同":
                agree += 1
            elif fb == "不认同":
                disagree += 1
            rows.append(
                {
                    "操作人名": meta["操作人名"],
                    "chatid": meta["chatid"],
                    "对话brief": meta["对话brief"],
                    "每轮对话问题描述": turn["每轮对话问题描述"],
                    "每轮对话反馈": fb,
                    "反馈描述": turn["反馈描述"],
                    "每轮对话生成的sql语句": turn["每轮对话生成的sql语句"],
                    "对话时间": turn["对话时间"],
                    "备注": meta["备注"],
                }
            )
            excel_row += 1
        end = excel_row - 1
        if end > start:
            merge_ranges.append((start, end))

    chat_count = len(by_chat)
    turn_count = len(rows)

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "对话导出"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="center")

    for row in rows:
        ws.append([row[h] for h in HEADERS])

    summary_row_idx = ws.max_row + 1
    summary_fill = PatternFill("solid", fgColor="FFF2CC")
    summary_font = Font(bold=True)
    summary_text = (
        f"汇总：总对话数={chat_count}；对话轮次数={turn_count}；"
        f"认同数={agree}；不认同数={disagree}"
    )
    ws.append([summary_text] + [""] * (len(HEADERS) - 1))
    ws.merge_cells(
        start_row=summary_row_idx,
        start_column=1,
        end_row=summary_row_idx,
        end_column=len(HEADERS),
    )
    summary_cell = ws.cell(summary_row_idx, 1)
    summary_cell.font = summary_font
    summary_cell.fill = summary_fill
    summary_cell.alignment = Alignment(vertical="center", wrap_text=True)
    for col in range(1, len(HEADERS) + 1):
        ws.cell(summary_row_idx, col).fill = summary_fill

    merge_cols = (1, 2, 3, 9)
    for start, end in merge_ranges:
        for col in merge_cols:
            ws.merge_cells(
                start_row=start,
                start_column=col,
                end_row=end,
                end_column=col,
            )

    widths = {
        "A": 14,
        "B": 12,
        "C": 36,
        "D": 48,
        "E": 12,
        "F": 24,
        "G": 80,
        "H": 22,
        "I": 40,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row_idx in range(2, summary_row_idx):
        for col in range(1, len(HEADERS) + 1):
            if col in (1, 2, 3, 5):
                align = Alignment(vertical="center", wrap_text=True, horizontal="center")
            else:
                align = Alignment(vertical="top", wrap_text=True)
            ws.cell(row_idx, col).alignment = align

    ws2 = wb.create_sheet("导出说明")
    ws2.append(["字段", "说明"])
    for item in [
        ("工作空间", f"oid={oid} name={workspace_name or ''}"),
        ("对话总数(列表)", str(len(chats))),
        ("排除 records 为空", str(skipped_empty)),
        ("排除 Admin", str(skipped_admin)),
        ("排除无SQL对话", str(skipped_no_sql)),
        ("导出对话数", str(chat_count)),
        ("导出行数(有SQL轮次)", str(turn_count)),
        ("认同数 / 不认同数", f"{agree} / {disagree}"),
        ("聚合", "按 chatid 分组；操作人名/chatid/对话brief/备注 合并单元格"),
        ("组内排序", "同一 chatid 内按对话时间升序"),
        ("SQL过滤", "整段对话无SQL则排除；有SQL时去掉无SQL轮次，提问写入备注"),
        ("反馈映射", "up→认同，down→不认同；反馈描述=feedback_comment"),
        ("汇总行", "总对话数 / 对话轮次数 / 认同数 / 不认同数"),
    ]:
        ws2.append(list(item))
    for col in range(1, 3):
        ws2.column_dimensions[get_column_letter(col)].width = 52

    buffer = io.BytesIO()
    wb.save(buffer)
    stats = {
        "oid": oid,
        "workspace_name": workspace_name,
        "chats_list": len(chats),
        "skipped_empty": skipped_empty,
        "skipped_admin": skipped_admin,
        "skipped_no_sql": skipped_no_sql,
        "chats_exported": chat_count,
        "rows": turn_count,
        "agree": agree,
        "disagree": disagree,
        "merged_chat_groups": len(merge_ranges),
    }
    return buffer.getvalue(), stats
